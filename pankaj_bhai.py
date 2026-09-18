import os
import time
import pandas as pd
from eqldata import generate_auth_token, get_1MARKET_DATA
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------
# 1. Configuration & Authentication
# ---------------------------------------------------------
username = 'support@teztrader.com'
password = '1XFBqsBTs0a6rimR'

instruments = [
    "NSEEQ:NIFTY_50",
]

# NSE Trading Holidays (Add 2025/2026 holidays covering your 6-month window)
NSE_HOLIDAYS = {
    "2025-10-02", "2025-10-21", "2025-10-22", "2025-11-05", "2025-12-25",
    "2026-01-15", "2026-01-26", "2026-03-03", "2026-03-26",
    "2026-03-31", "2026-04-03", "2026-04-14", "2026-05-01",
    "2026-05-28", "2026-06-26", "2026-09-14", "2026-10-02",
    "2026-10-20", "2026-11-10", "2026-11-24", "2026-12-25"
}

# ---------------------------------------------------------
# 2. Date and API Helper Functions
# ---------------------------------------------------------
def get_trading_days_past_6_months():
    """Generates all valid NSE trading days over the last 182 days (~6 months)."""
    end_date = datetime.now()
    start_date = end_date - timedelta(days=182)
    
    valid_dates = []
    current_date = start_date

    while current_date <= end_date:
        date_str = current_date.strftime("%Y-%m-%d")
        is_weekend = current_date.weekday() >= 5
        is_holiday = date_str in NSE_HOLIDAYS

        if not is_weekend and not is_holiday:
            valid_dates.append(date_str)

        current_date += timedelta(days=1)

    return valid_dates

def safe_fetch_market_data(token, instrument_list, target_date):
    """Rate-limit wrapper to fetch 1-min market data."""
    max_retries = 3
    for attempt in range(max_retries):
        try:
            result = get_1MARKET_DATA(token, instrument_list, target_date)
            
            if isinstance(result, list) and len(result) > 0 and isinstance(result[0], list):
                return result

            result_str = str(result).lower()
            if "too many requests" in result_str or "429" in result_str:
                print(f" [!] Intraday Rate Limited. Waiting 30s (Attempt {attempt + 1}/{max_retries})...")
                time.sleep(30)
                continue

            return result
        except Exception as e:
            error_str = str(e).lower()
            if "too many requests" in error_str or "429" in error_str:
                print(f" [!] Intraday Rate Limited (Exception). Waiting 30s...")
                time.sleep(30)
                continue
            else:
                print(f" [X] Unexpected Intraday API Error: {e}")
                return None
    return None

# ---------------------------------------------------------
# 3. Data Processing & Excel Export
# ---------------------------------------------------------
def process_and_export_data(raw_data_list):
    if not raw_data_list:
        print("No valid intraday data received to process.")
        return

    print("Building 1-minute DataFrame...")
    columns = ['Symbol', 'Timestamp', 'Open', 'High', 'Low', 'Close', 'Volume', 'Extra']
    df_1m = pd.DataFrame(raw_data_list, columns=columns)

    # Clean timestamp and format
    df_1m['Timestamp'] = df_1m['Timestamp'].astype(str).str.replace(' ', '')
    df_1m['Timestamp'] = pd.to_datetime(df_1m['Timestamp'], format='%Y-%m-%d%H:%M:%S', errors='coerce')
    
    df_1m.dropna(subset=['Timestamp'], inplace=True)
    if 'Extra' in df_1m.columns:
        df_1m.drop(columns=['Extra'], inplace=True)

    cols_to_numeric = ['Open', 'High', 'Low', 'Close', 'Volume']
    df_1m[cols_to_numeric] = df_1m[cols_to_numeric].apply(pd.to_numeric)

    df_1m.set_index('Timestamp', inplace=True)
    df_1m.sort_index(inplace=True)

    print("Resampling 1-minute data into 3-minute candles...")
    df_3m = df_1m.groupby('Symbol').resample('3min').agg({
        'Open': 'first',
        'High': 'max',
        'Low': 'min',
        'Close': 'last',
        'Volume': 'sum'
    }).dropna()

    df_1m.reset_index(inplace=True)
    df_3m.reset_index(inplace=True)
    
    cols_order = ['Symbol', 'Timestamp', 'Open', 'High', 'Low', 'Close', 'Volume']
    df_1m = df_1m[cols_order]
    df_3m = df_3m[cols_order]

    export_to_excel(df_1m, df_3m)

def export_to_excel(df_1m, df_3m):
    filename = f"NIFTY50_6Months_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    file_path = os.path.join(os.getcwd(), filename)
    
    print(f"Writing {len(df_1m):,} 1-min rows and {len(df_3m):,} 3-min rows to Excel...")
    wb = Workbook()
    
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
        
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    # Pre-defined optimal column widths for fast rendering with ~46k rows
    col_widths = {
        'Symbol': 18,
        'Timestamp': 22,
        'Open': 12,
        'High': 12,
        'Low': 12,
        'Close': 12,
        'Volume': 14
    }

    def write_sheet(ws_name, df):
        ws = wb.create_sheet(title=ws_name)
        headers = df.columns.tolist()
        ws.append(headers)
        
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(header, 15)
            
        df_export = df.copy()
        df_export['Timestamp'] = df_export['Timestamp'].dt.strftime('%Y-%m-%d %H:%M:%S')
        
        for r in df_export.to_numpy():
            ws.append(r.tolist())

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    write_sheet("1 Min Data", df_1m)
    write_sheet("3 Min Data", df_3m)

    wb.save(file_path)
    print(f"\n[+] Successfully saved Excel file: {file_path}")

# ---------------------------------------------------------
# 4. Main Execution Block
# ---------------------------------------------------------
if __name__ == "__main__":
    auth_token = generate_auth_token(username, password)
    
    if auth_token:
        print("Authentication Successful.\n")
        
        all_combined_data = []
        target_dates = get_trading_days_past_6_months()
        
        print(f"Targeting {len(target_dates)} trading days from {target_dates[0]} to {target_dates[-1]}.\n")

        for idx, target_date in enumerate(target_dates, start=1):
            print(f"[{idx}/{len(target_dates)}] Fetching: {target_date}...")
            result = safe_fetch_market_data(auth_token, instruments, target_date)

            def extract_ticks(lst):
                for item in lst:
                    if isinstance(item, list):
                        if len(item) >= 7 and isinstance(item[0], str) and item[0] != 'No Data':
                            all_combined_data.append(item)
                        else:
                            extract_ticks(item)

            if result and isinstance(result, list):
                extract_ticks(result)
            else:
                print(f"   -> No data returned for {target_date}.")

            # Small polite pause to prevent triggering 429 rate limits
            time.sleep(0.3)

        print(f"\nTotal 1-min raw ticks loaded: {len(all_combined_data):,}")
        
        if all_combined_data:
            process_and_export_data(all_combined_data)
        else:
            print("No data collected. Exiting.")